# -*- coding: utf-8 -*-
"""Build the blitz disti pack so every post is ONE clean copy-paste block.

The first version used markdown blockquotes and hard-wrapped at 80 columns. Both are
wrong for this job: the "> " prefixes have to be stripped line by line, and a hard wrap
mid-sentence becomes a real line break the moment it is pasted into LinkedIn. Post copy
here is stored unwrapped, with line breaks ONLY where one is intended."""
import os, re
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

POSTS = [
 dict(n=1, title="The AEO Buyer's Standard",
  image="public/img/buyer-series/aeo-buyers-standard-hero-1600x900.jpg",
  url="https://aeoanalyzers.com/aeo-buyers-standard",
  why="Lead with the asset, not the ad. Most linkable, most forwardable, most likely to be cited.",
  li=["We published the standard we think buyers should use to judge AEO software — including us.",
      "Seven requirements. Observe, diagnose, remediate, verify, receipts, scale, economics. Thirteen questions to take into a vendor call.",
      "It names no vendor and ranks nothing. Then it scores us against all seven, in full, including the one where we come off worst: scale. We have narrower engine coverage than the big platforms and no multi-brand or SSO tooling, and a portfolio operator should look elsewhere. That's on the page.",
      "The line I'd keep if you keep one: AEO metrics should have receipts. If a dashboard says your visibility is 42%, you should be able to ask 42% of what — which prompts, which engines, which dates, how many runs — and read the answer behind the number.",
      "Take it into a procurement meeting. Use it on us. A standard that only works when it flatters the company that wrote it isn't a standard.",
      "If a requirement is missing from the seven, tell me — I'll add it and credit whoever suggested it.",
      "https://aeoanalyzers.com/aeo-buyers-standard",
      "#AEO #AISearch #GenerativeAI #B2BMarketing #Procurement"],
  x=["We published the standard we think buyers should use to judge AEO software — and scored ourselves against it, including where we come off worst.",
     "7 requirements. 13 questions. Names no vendor, ranks nothing.",
     "AEO metrics should have receipts.",
     "https://aeoanalyzers.com/aeo-buyers-standard"]),

 dict(n=2, title="Why AI still doesn't mention your business",
  image="public/blog/why-ai-doesnt-mention-you/img/primer-og-1200x630.jpg",
  url="https://aeoanalyzers.com/blog/why-ai-doesnt-mention-you",
  why="The primer. Widest audience, least commercial, strongest numbers.",
  li=["You did the work the guides told you to do. Nothing changed.",
      "Here's some of what we measured and what the published research says:",
      "→ llms.txt goes unrequested on 97% of the sites that have one\n→ six of seven AI platforms cannot read your schema\n→ the assistants read different indexes, so being in one doesn't put you in another\n→ an AI sees roughly the first 200 characters after your h1, and ignores your meta description entirely",
      "That last one is the cheapest fix in this whole field and almost nobody does it.",
      "The piece includes the free self-checks — no tool required — and two of our own failures, because a primer that only lists other people's mistakes isn't worth reading.",
      "Run the free self-checks on your own site and tell me what you find. I'll answer anything in the comments.",
      "https://aeoanalyzers.com/blog/why-ai-doesnt-mention-you",
      "#AEO #AISearch #SEO #GenerativeEngineOptimization #ContentStrategy"],
  x=["You did what the AEO guides said. Nothing changed.",
     "llms.txt unrequested on 97% of sites that have one. 6 of 7 platforms can't read your schema. An AI sees ~200 characters after your h1.",
     "The free self-checks, and two of our own failures:",
     "https://aeoanalyzers.com/blog/why-ai-doesnt-mention-you"]),

 dict(n=3, title="What should an AEO tool actually do?",
  image="public/img/buyer-series/what-should-an-aeo-tool-do-hero-1600x900.jpg",
  url="https://aeoanalyzers.com/what-should-an-aeo-tool-do",
  why="Defines the vocabulary the rest of the series uses. Teaching, not selling.",
  li=["\"AI visibility\" is at least six different outcomes, and the fixes have almost nothing in common.",
      "The engine knows you exist. It mentions you. It links to you. It cites you as a source. It describes you accurately. It recommends you when a buyer never said your name.",
      "Those are not interchangeable wins, and a single blended score hides the difference. A business can be perfectly findable by name and never recommended to anyone — which isn't a weaker version of the same problem. It's a different problem with a different fix.",
      "So before comparing tools, work out which of these four you're actually buying:",
      "Monitoring asks what happened. Diagnosis asks why. Remediation asks what to change. Verification asks whether the change worked.",
      "Most products do more than one. Almost none do all four equally well, and the price rarely tells you which.",
      "Which of those four are you actually paying for? Genuinely curious what people find they bought.",
      "https://aeoanalyzers.com/what-should-an-aeo-tool-do",
      "#AEO #GEO #AISearch #MarTech #SEO"],
  x=["\"AI visibility\" is six different outcomes with almost nothing in common.",
     "A mention is not a citation. A citation is not a recommendation. Measurement is not diagnosis.",
     "https://aeoanalyzers.com/what-should-an-aeo-tool-do"]),

 dict(n=4, title="Best AEO tools in 2026",
  image="public/img/buyer-series/best-aeo-tools-hero-1600x900.jpg",
  url="https://aeoanalyzers.com/best-aeo-tools",
  why="The commercial page. Lands better after the standard exists. Always disclose we make one of these tools.",
  li=["We wrote the comparison of our own category, and we don't name a winner.",
      "Profound, Peec AI, Otterly AI, Scrunch AI, AthenaHQ, Rankscale, SE Ranking — and us. Compared on seven purchasing requirements rather than feature counts.",
      "Three things we made ourselves do:",
      "→ Every competitor fact carries the date it was verified and links to that vendor's own published page. Nothing from another comparison site, nothing from memory.\n→ We say plainly where each one is the better choice. Need breadth across many engines and regions? That's Rankscale. Want SEO and GEO in one suite? SE Ranking. Agent infrastructure? Scrunch.\n→ We say where we're the weaker fit, which is any portfolio-scale operation.",
      "We make one of these tools and the page says so in the first paragraph. Treat it as interested testimony and check it — every claim is dated and sourced so that you can.",
      "If you think we've got a competitor wrong, tell me. Every fact is dated and I'll correct it with a note saying what changed.",
      "https://aeoanalyzers.com/best-aeo-tools",
      "#AEO #AISearch #MarTech #B2BSaaS #SEO"],
  x=["We wrote the comparison of our own category and named no winner.",
     "8 tools, 7 purchasing requirements, every competitor fact dated and linked to that vendor's own page.",
     "Including where we're the weaker choice.",
     "https://aeoanalyzers.com/best-aeo-tools"]),

 dict(n=5, title="The AEO Buyer's Guide — three short reads",
  image="public/img/buyer-series/aeo-buyers-guide-hero-1600x900.jpg",
  url="https://aeoanalyzers.com/aeo-buyers-guide",
  why="The catch-all for anyone who missed 1 to 4. Note: this image is a known mismatch until its replacement is generated.",
  li=["Six minutes, three short reads, and you'll know how to buy AEO software.",
      "1. What an AEO tool should actually do — why \"AI visibility\" is six different outcomes.\n2. What to require before you buy — seven purchasing requirements and the questions that go with them.\n3. How to compare the market — why there's no single best tool until you define the job.",
      "Each one is about two minutes and links to the full guide behind it, so you can go as deep as you need and stop when you have what you came for.",
      "Written by a vendor in the category, which the page says at the top. The standard inside it works against us as readily as anyone else — that's the point of publishing it.",
      "Six minutes. If it saves you one bad vendor call it has paid for itself.",
      "https://aeoanalyzers.com/aeo-buyers-guide",
      "#AEO #AISearch #Procurement #B2BMarketing #GEO"],
  x=["How to buy AEO software, in six minutes.",
     "Three two-minute reads: what a tool should do, what to require before you buy, how to compare the market. Each links to the full guide behind it.",
     "https://aeoanalyzers.com/aeo-buyers-guide"]),

 dict(n=6, title="How AEO Analyzers works",
  image="public/og-blog-how-it-works.png",
  url="https://aeoanalyzers.com/blog/how-it-works",
  why="Product mechanics. Last, because it's the only one that's about us.",
  li=["The whole pipeline, drawn. Six diagrams, no algorithms.",
      "We read your site. We ask four answer engines your buyers' questions, several times each, with web search on. We keep search-grounded answers separate from ones the model produced from memory, because blending them moves a score invisibly.",
      "Then we score three layers separately rather than blending them into one number: whether an engine finds you when asked by name, whether what it says about you is accurate, and whether it recommends you to a buyer who never said your name.",
      "Every number is backed by a stored transcript. Open any figure, read the answer behind it, and run the same question yourself to check we're telling you the truth.",
      "That last part is the whole design. A number you can't reproduce is a number you can't dispute.",
      "Run any of those queries yourself and tell me if you get something different. That is the whole promise.",
      "https://aeoanalyzers.com/blog/how-it-works",
      "#AEO #AISearch #Measurement #GenerativeAI #SEO"],
  x=["The whole pipeline, drawn. Six diagrams, no algorithms.",
     "Four engines, buyers' questions, several runs each. Search-grounded answers kept separate from memory. Three layers scored separately, every number backed by a stored transcript you can re-run.",
     "https://aeoanalyzers.com/blog/how-it-works"]),
]


# --------------------------------------------------------------- LENGTH GUARD
def x_length(parts):
    """X counts every URL as 23 characters, whatever its real length."""
    total = 0
    for i, b in enumerate(parts):
        total += 23 if b.startswith('http') else len(b)
        if i < len(parts) - 1:
            total += 2   # blank line between paragraphs
    return total

for _p in POSTS:
    _n = x_length(_p['x'])
    assert _n <= 280, f"X post {_p['n']} is {_n} chars — over the 280 limit"

RULES = ("Every post carries its link in the body — the autoposter has no comment API, so there is no "
         "other placement, and the copy is written to carry the URL naturally. No cost or expense figure "
         "appears anywhere. Every number is reproducible from a stored transcript or a published study; "
         "nothing is invented, not even as an example. No claim that other tools \"only give you a score\".")
RHYTHM = ("All six can go the same day — async audiences, 20 to 100 readers a post, so order matters more "
          "than spacing. If you would rather stagger: 1 and 2 on day one about three hours apart, 3 and 4 "
          "on day two, 5 and 6 on day three. Post to LinkedIn first each time, then X about thirty minutes "
          "later. The LinkedIn post is canonical; the X post is a pointer.")

# ------------------------------------------------------------------ MARKDOWN
def markdown():
    L=["# BLITZ DISTI PACK — AEO buyer series","",
       "Six posts. Each block below is ready to select and copy in one go — no prefixes, no "
       "hard wraps, nothing to strip.","",
       "**Standing rules:** "+RULES,"","**Rhythm:** "+RHYTHM,"","---","","## Publish order","",
       "| # | Post | Image to attach |","|---|---|---|"]
    for p in POSTS:
        L.append(f"| {p['n']} | {p['title']} | `{os.path.basename(p['image'])}` |")
    L += ["","---",""]
    for p in POSTS:
        L += [f"## {p['n']} · {p['title']}","",
              f"**Why here:** {p['why']}","",
              f"**Image:** `{p['image']}`","",
              f"**Link:** {p['url']}","",
              "### LinkedIn — copy from here","",
              "\n\n".join(p['li']),"",
              "### X — copy from here","",
              "\n\n".join(p['x']),"","---",""]
    return "\n".join(L)

# ---------------------------------------------------------------------- DOCX
def shade(par, hexcolor):
    pPr=par._p.get_or_add_pPr(); s=OxmlElement('w:shd')
    s.set(qn('w:val'),'clear'); s.set(qn('w:fill'),hexcolor); pPr.append(s)

def docx(path):
    d=Document()
    for s in d.sections:
        s.left_margin=s.right_margin=Inches(0.9)
    n=d.styles['Normal']; n.font.name='Calibri'; n.font.size=Pt(11)
    d.add_heading('BLITZ DISTI PACK — AEO buyer series',0)
    i=d.add_paragraph('Each grey block is one post. Select the whole block and copy — no prefixes, no line-break artifacts.')
    i.runs[0].italic=True
    d.add_heading('Standing rules',2); d.add_paragraph(RULES)
    d.add_heading('Rhythm',2); d.add_paragraph(RHYTHM)
    d.add_heading('Publish order',2)
    t=d.add_table(rows=1,cols=3); t.style='Light Grid Accent 1'
    for c,h in zip(t.rows[0].cells,['#','Post','Image to attach']):
        c.text=h
        for par in c.paragraphs:
            for r in par.runs: r.bold=True
    for p in POSTS:
        c=t.add_row().cells
        c[0].text=str(p['n']); c[1].text=p['title']; c[2].text=os.path.basename(p['image'])

    for p in POSTS:
        d.add_page_break()
        d.add_heading(f"{p['n']} · {p['title']}",1)
        for label,val in [('Why here',p['why']),('Image',p['image']),('Link',p['url'])]:
            par=d.add_paragraph(); par.add_run(label+': ').bold=True; par.add_run(val)
        for kind,body in [('LinkedIn',p['li']),('X',p['x'])]:
            h=d.add_heading(f'{kind} — select this block and copy',2)
            for j,block in enumerate(body):
                par=d.add_paragraph(block)
                par.paragraph_format.left_indent=Inches(0.12)
                par.paragraph_format.right_indent=Inches(0.12)
                par.paragraph_format.space_after=Pt(10 if j<len(body)-1 else 16)
                shade(par,'F2F4F5')
    d.save(path)

# Guarded so the 5-6-7 pack can import POSTS and reuse the SAME post text rather than
# retyping it — one source of truth for the copy, two documents built from it.
if __name__ == '__main__':
    open('docs/launch/BLITZ-DISTI-PACK-aeo-buyer-series.md','w',encoding='utf-8').write(markdown())
    docx('docs/launch/BLITZ-DISTI-PACK-aeo-buyer-series.docx')
    print('rebuilt')
